# -*- coding: utf-8 -*-
"""Сборка отчёта за 2 семестр по образцу оформления (исправленный отчёт 1 семестра).

    python build_report_sem2.py <root>

Шаблон — src/Отчет 1 сем (исправленный, образец оформления).docx: из него берутся титульный лист,
стили, содержание, заголовки, абзацы, таблица, список литературы и оформление статьи в приложении.
Текст — report_content_sem2.py. Сборка повторяется (docx → PDF через LibreOffice → разбор PDF), пока
не установятся номера страниц в содержании и разбиение таблиц по страницам: таблица, не помещающаяся
на странице, либо переносится целиком, либо (длинная) делится с подписью «Продолжение таблицы N».
"""
import copy
import json
import pathlib
import re
import subprocess
import sys
import time

import docx
import latex2mathml.converter
import mathml2omml
import pymupdf
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm

sys.stdout.reconfigure(encoding="utf-8")
root = pathlib.Path(sys.argv[1])
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import report_content_sem2 as C  # noqa: E402

SRC = root / "src" / "Отчет 1 сем (исправленный, образец оформления).docx"
OUT = root / "report" / "Отчет_Филадельфов_09-535_2_сем.docx"
PDF = OUT.with_suffix(".pdf")
FIG = root / "report" / "sem2" / "fig"
SOFFICE = r"C:\Program Files\LibreOffice\program\soffice.exe"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

# элементы образца: номер в теле документа
PROTO = {"h1": 40, "p": 41, "ch": 47, "tcap": 77, "tbl": 78, "tcont": 80, "ref": 84,
         "atitle": 95, "aauth": 96, "asup": 97, "aaff": 98, "amail": 99, "ap": 100}


def W(tag):
    return qn("w:" + tag)


# ------------------------------------------------------------------ работа с XML абзацев
def ppr_of(el):
    ppr = el.find(W("pPr"))
    if ppr is None:
        ppr = el.makeelement(W("pPr"), {})
        el.insert(0, ppr)
    return ppr


def ppr_set(el, tag, attrs=None, before=("rPr",)):
    """Заменить (или добавить) свойство абзаца, вставив его перед rPr."""
    ppr = ppr_of(el)
    for old in ppr.findall(W(tag)):
        ppr.remove(old)
    new = ppr.makeelement(W(tag), {W(k): v for k, v in (attrs or {}).items()})
    anchor = next((ppr.find(W(t)) for t in before if ppr.find(W(t)) is not None), None)
    (anchor.addprevious(new) if anchor is not None else ppr.append(new))
    return el


def set_text(el, text, bold=None):
    """Оставить первый текстовый run абзаца с его форматированием и записать в него текст."""
    runs = [r for r in el.iter(W("r")) if r.find(W("t")) is not None]
    keep = runs[0]
    for child in list(el):
        if child.tag != W("pPr") and child is not keep and keep not in list(child.iter()):
            el.remove(child)
    if keep.getparent() is not el:
        el.append(keep)
    for c in list(keep):
        if c.tag != W("rPr"):
            keep.remove(c)
    t = keep.makeelement(W("t"), {})
    t.text = text
    t.set(XML_SPACE, "preserve")
    keep.append(t)
    if bold is not None:
        rpr = keep.find(W("rPr"))
        for b in rpr.findall(W("b")) + rpr.findall(W("bCs")):
            rpr.remove(b)
        if bold:
            rpr.append(rpr.makeelement(W("b"), {}))
    return el


def strip_marks(el):
    """Закладки, отметки правописания и «последний разрыв страницы» из образца не нужны."""
    for tag in ("bookmarkStart", "bookmarkEnd", "proofErr", "lastRenderedPageBreak"):
        for x in list(el.iter(W(tag))):
            x.getparent().remove(x)
    return el


def keep_next(el):
    return ppr_set(el, "keepNext", before=("pageBreakBefore", "snapToGrid", "numPr", "spacing", "ind", "jc", "rPr"))


def page_break_before(el):
    return ppr_set(el, "pageBreakBefore", before=("snapToGrid", "numPr", "spacing", "ind", "jc", "rPr"))


# ------------------------------------------------------------------ формулы (редактор формул Word, OMML)
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
MATH = re.compile(r"\$(.+?)\$")


def omath(tex):
    """LaTeX -> элемент m:oMath; у каждого фрагмента шрифт Cambria Math 14 pt (так формулы выглядят в Word)."""
    xml = mathml2omml.convert(latex2mathml.converter.convert(tex))
    el = parse_xml(f'<p xmlns:m="{M_NS}" xmlns:w="{W_NS}">{xml}</p>')[0]
    for r in el.iter(f"{{{M_NS}}}r"):
        rpr = parse_xml(f'<w:rPr xmlns:w="{W_NS}"><w:rFonts w:ascii="Cambria Math" w:hAnsi="Cambria Math"/>'
                        '<w:sz w:val="28"/><w:szCs w:val="28"/></w:rPr>')
        mrpr = r.find(f"{{{M_NS}}}rPr")
        (mrpr.addnext(rpr) if mrpr is not None else r.insert(0, rpr))
    return el


def fill_rich(el, text):
    """Текст абзаца с формулами в $...$: обычный текст — run по образцу, формулы — m:oMath."""
    runs = [r for r in el.iter(W("r")) if r.find(W("t")) is not None]
    rpr = copy.deepcopy(runs[0].find(W("rPr")))
    for child in list(el):
        if child.tag != W("pPr"):
            el.remove(child)
    pos = 0
    for m in list(MATH.finditer(text)) + [None]:
        chunk = text[pos:m.start()] if m else text[pos:]
        if chunk:
            r = el.makeelement(W("r"), {})
            r.append(copy.deepcopy(rpr))
            t = r.makeelement(W("t"), {})
            t.text = chunk
            t.set(XML_SPACE, "preserve")
            r.append(t)
            el.append(r)
        if m:
            el.append(omath(m.group(1)))
            pos = m.end()
    return el


# ------------------------------------------------------------------ нумерация источников по порядку упоминания
CITE = re.compile(r"\[(\d+(?:\s*[,–-]\s*\d+)*)\]")


def cite_nums(group):
    nums = []
    for part in group.split(","):
        part = part.strip()
        if re.fullmatch(r"\d+\s*[–-]\s*\d+", part):
            a, z = map(int, re.split(r"\s*[–-]\s*", part))
            nums += list(range(a, z + 1))
        else:
            nums.append(int(part))
    return nums


def cite_fmt(nums):
    """[3, 4, 5, 8] -> «3–5, 8»."""
    nums, out, i = sorted(set(nums)), [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        out.append(f"{nums[i]}–{nums[j]}" if j - i >= 2 else ", ".join(map(str, nums[i:j + 1])))
        i = j + 1
    return ", ".join(out)


_order = []
for b in C.BLOCKS:
    if b[0] in ("p", "where"):
        for m in CITE.finditer(b[1]):
            _order += [n for n in cite_nums(m.group(1)) if n not in _order]
_refs = next(b[1] for b in C.BLOCKS if b[0] == "refs")
_order += [n for n in range(1, len(_refs) + 1) if n not in _order]
RENUM = {old: new for new, old in enumerate(_order, 1)}


def renum(text):
    return CITE.sub(lambda m: "[" + cite_fmt([RENUM[n] for n in cite_nums(m.group(1))]) + "]", text)


# ------------------------------------------------------------------ содержание: заголовки из блоков
def toc_entries():
    out, ch = [], 0
    for b in C.BLOCKS:
        if b[0] == "h1":
            out.append((b[1], 1))
        elif b[0] == "ch":
            ch += 1
            out.append((f"{ch}. {b[1]}", 2))
        elif b[0] == "sub":
            out.append((b[1], 3))
    return out


TABLES = [(int(re.match(r"Таблица (\d+)", b[2]).group(1)), b) for b in C.BLOCKS if b[0] == "table"]


def table_no(b):
    return int(re.match(r"Таблица (\d+)", b[2]).group(1))


def arrange(splits):
    """Порядок блоков с учётом решений: у «отложенной» таблицы следующие за ней абзацы идут перед ней,
    чтобы заполнить низ страницы, а таблица начинается на следующей (ГОСТ допускает таблицу после ссылки)."""
    out, i, B = [], 0, C.BLOCKS
    while i < len(B):
        b = B[i]
        if b[0] == "table" and splits.get(table_no(b), {}).get("defer"):
            j = i + 1
            while j < len(B) and B[j][0] == "p":
                j += 1
            out += list(B[i + 1:j]) + [b]
            i = j
        else:
            out.append(b)
            i += 1
    return out


def can_defer(number):
    B = C.BLOCKS
    i = next(i for i, b in enumerate(B) if b[0] == "table" and table_no(b) == number)
    return i + 1 < len(B) and B[i + 1][0] == "p"
LONG = {n for n, b in TABLES if "компетенций" in b[2] or len(b[4]) > 10}  # такие таблицы делятся


# ------------------------------------------------------------------ сборка
def build(toc_pages, splits):
    D = docx.Document(SRC)
    body = D.element.body
    kids = list(body.iterchildren())
    T = {k: strip_marks(copy.deepcopy(kids[i])) for k, i in PROTO.items()}
    sdt, sect = kids[38], kids[-1]
    assert sect.tag == W("sectPr") and sdt.tag == W("sdt")

    # титульный лист образца: только год
    for el in kids[:33]:
        for t in el.iter(W("t")):
            if t.text and "2025" in t.text:
                t.text = t.text.replace("2025", "2026")
    for el in kids[33:38] + kids[39:-1]:
        body.remove(el)

    def add(el):
        sect.addprevious(el)
        return el

    def para(key, text, bold=None):
        return add(set_text(copy.deepcopy(T[key]), text, bold))

    # ---- содержание
    content = sdt.find(W("sdtContent"))
    tps = content.findall(W("p"))
    top_ppr = copy.deepcopy(tps[1].find(W("pPr")))
    sub_ppr = copy.deepcopy(tps[2].find(W("pPr")))

    def text_rpr(par):
        run = next(r for r in par.iter(W("r")) if r.find(W("t")) is not None)
        rpr = copy.deepcopy(run.find(W("rPr")))
        for x in rpr.findall(W("rStyle")):
            rpr.remove(x)
        return rpr
    top_rpr, sub_rpr = text_rpr(tps[1]), text_rpr(tps[2])
    for p in tps[1:]:
        content.remove(p)
    for title, level in toc_entries():
        p = content.makeelement(W("p"), {})
        ppr = copy.deepcopy(top_ppr if level == 1 else sub_ppr)
        p.append(ppr)
        if level > 1:  # главы и подразделы — без отступов, прижаты влево (замечание преподавателя)
            ppr_set(p, "ind", {"left": "0", "firstLine": "0"}, before=("jc", "rPr"))
        for kind, val in (("t", title), ("tab", None), ("t", str(toc_pages.get(title, "")))):
            r = p.makeelement(W("r"), {})
            r.append(copy.deepcopy(top_rpr if level == 1 else sub_rpr))
            if kind == "tab":
                r.append(r.makeelement(W("tab"), {}))
            else:
                t = r.makeelement(W("t"), {})
                t.text = val
                t.set(XML_SPACE, "preserve")
                r.append(t)
            p.append(r)
        content.append(p)

    # ---- таблица по образцу: шрифт 12, рамки, шапка по центру
    def make_table(header, rows, widths_cm, align, repeat_header):
        tbl = copy.deepcopy(T["tbl"])
        trs = tbl.findall(W("tr"))
        proto_h, proto_d = trs[0], trs[1]
        for tr in trs:
            tbl.remove(tr)
        n = len(header)
        dxa = [int(w * 567) for w in widths_cm]
        grid = tbl.find(W("tblGrid"))
        for g in list(grid):
            grid.remove(g)
        for w in dxa:
            grid.append(grid.makeelement(W("gridCol"), {W("w"): str(w)}))
        tw = tbl.find(W("tblPr")).find(W("tblW"))
        tw.set(W("w"), str(sum(dxa)))
        tw.set(W("type"), "dxa")
        jcs = {"l": "left", "c": "center", "r": "right", "j": "both"}

        def row(proto, values, head):
            tr = copy.deepcopy(proto)
            trpr = tr.find(W("trPr"))
            if trpr is None:
                trpr = tr.makeelement(W("trPr"), {})
                tr.insert(0, trpr)
            for x in trpr.findall(W("tblHeader")) + trpr.findall(W("cantSplit")):
                trpr.remove(x)
            trpr.insert(0, trpr.makeelement(W("cantSplit"), {}))
            if head and repeat_header:
                trpr.append(trpr.makeelement(W("tblHeader"), {}))
            tcs = tr.findall(W("tc"))
            while len(tcs) < n:
                tcs[-1].addnext(copy.deepcopy(tcs[-1]))
                tcs = tr.findall(W("tc"))
            for extra in tcs[n:]:
                tr.remove(extra)
            for tc, v, w, a in zip(tr.findall(W("tc")), values, dxa, align):
                tcpr = tc.find(W("tcPr"))
                tcpr.find(W("tcW")).set(W("w"), str(w))
                mar = tcpr.find(W("tcMar"))  # небольшие поля слева и справа, чтобы текст не касался рамки
                if mar is not None:
                    for side in ("left", "right"):
                        e = mar.find(W(side))
                        if e is not None:
                            e.set(W("w"), "57")
                ps = tc.findall(W("p"))
                for extra in ps[1:]:
                    tc.remove(extra)
                strip_marks(ps[0])
                set_text(ps[0], v)
                ppr_set(ps[0], "jc", {"val": "center" if head else jcs[a]})
            tbl.append(tr)

        row(proto_h, header, True)
        for r in rows:
            row(proto_d, r, False)
        return tbl

    def table(number, caption, header, rows, widths, align):
        align = align.ljust(len(header), "l")
        dec = splits.get(number, {})
        move, k = dec.get("move", False), dec.get("k")
        cap = keep_next(para("tcap", caption))
        if move:  # таблица начинается с новой страницы
            page_break_before(cap)
        parts = [rows] if not k else [rows[:k], rows[k:]]
        for i, part in enumerate(parts):
            if i == 1:
                cont = keep_next(para("tcont", f"Продолжение таблицы {number}"))
                page_break_before(cont)
            add(make_table(header, part, widths, align, repeat_header=not k))

    def figure(fname, width_cm, caption):
        el = copy.deepcopy(T["p"])
        for r in list(el.iter(W("r"))):
            r.getparent().remove(r)
        ppr_set(el, "spacing", {"before": "160", "after": "0", "line": "240", "lineRule": "auto"})
        ppr_set(el, "ind", {"left": "0", "right": "0", "firstLine": "0"})
        ppr_set(el, "jc", {"val": "center"})
        add(keep_next(el))
        docx.text.paragraph.Paragraph(el, D._body).add_run().add_picture(str(FIG / fname), width=Cm(width_cm))
        cap = para("p", caption)
        ppr_set(cap, "ind", {"left": "0", "right": "0", "firstLine": "0"})
        ppr_set(cap, "jc", {"val": "center"})

    def code_block(text):
        for line in text.splitlines():
            lead = len(line) - len(line.lstrip(" "))
            el = para("p", "\u00a0" * lead + line.lstrip(" ") if line.strip() else "\u00a0")
            ppr_set(el, "spacing", {"before": "0", "after": "0", "line": "240", "lineRule": "auto"})
            ppr_set(el, "ind", {"left": "0", "right": "0", "firstLine": "0"})
            ppr_set(el, "jc", {"val": "left"})
            for rpr in el.iter(W("rPr")):
                f = rpr.find(W("rFonts"))
                if f is not None:
                    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
                        f.set(W(a), "Courier New")
                for sz in rpr.findall(W("sz")) + rpr.findall(W("szCs")):
                    sz.set(W("val"), "20")

    # ---- текст по блокам
    ch_no, eq_no, new_page = 0, 0, False
    for b in arrange(splits):
        kind = b[0]
        if kind == "pb":
            new_page = True
            continue
        mark = len(body)
        if kind == "h1":
            para("h1", b[1])
        elif kind == "ch":
            ch_no += 1
            keep_next(para("ch", f"{ch_no}. {b[1]}"))
        elif kind == "sub":
            keep_next(para("ch", b[1]))
        elif kind == "p":
            text = renum(b[1])
            (add(fill_rich(copy.deepcopy(T["p"]), text)) if "$" in text else para("p", text))
        elif kind == "where":  # пояснение к формуле: «где …» без абзацного отступа
            el = add(fill_rich(copy.deepcopy(T["p"]), renum(b[1])))
            ppr_set(el, "ind", {"left": "0", "firstLine": "0"})
        elif kind == "eq":  # формула отдельной строкой: по центру, номер в скобках справа
            eq_no += 1
            el = copy.deepcopy(T["p"])
            rpr = copy.deepcopy(next(r for r in el.iter(W("r")) if r.find(W("t")) is not None).find(W("rPr")))
            for child in list(el):
                if child.tag != W("pPr"):
                    el.remove(child)
            ppr_set(el, "tabs")
            tabs = ppr_of(el).find(W("tabs"))
            for val, pos in (("center", "4677"), ("right", "9354")):
                tabs.append(tabs.makeelement(W("tab"), {W("val"): val, W("pos"): pos}))
            ppr_set(el, "spacing", {"before": "60", "after": "60", "line": "240", "lineRule": "auto"})
            ppr_set(el, "ind", {"left": "0", "right": "0", "firstLine": "0"})
            ppr_set(el, "jc", {"val": "left"})
            for part in ("tab", "math", "tab", "num"):
                if part == "math":
                    el.append(omath(b[1]))
                    continue
                r = el.makeelement(W("r"), {})
                r.append(copy.deepcopy(rpr))
                if part == "tab":
                    r.append(r.makeelement(W("tab"), {}))
                else:
                    t = r.makeelement(W("t"), {})
                    t.text = f"({eq_no})"
                    r.append(t)
                el.append(r)
            add(el)
        elif kind == "table":
            _, _tpl, cap, head, rows, widths, align = b
            table(int(re.match(r"Таблица (\d+)", cap).group(1)), cap, head, rows, widths, align)
        elif kind == "fig":
            figure(b[1], b[2], b[3])
        elif kind == "refs":
            for old in _order:
                para("ref", b[1][old - 1])
        elif kind == "art":  # статья: заголовок, автор, руководитель, организация, почта — как в образце
            _, title, paras = b
            para("atitle", title)
            for k in ("aauth", "asup", "aaff", "amail"):
                add(copy.deepcopy(T[k]))
            for t in paras:
                para("ap", t)
        elif kind == "code":
            code_block(b[1])
        else:
            raise ValueError(kind)
        if new_page:
            first = list(body)[mark - 1]
            page_break_before(first)
            new_page = False

    # номера страниц — Times New Roman 12
    for sec in D.sections:
        for ft in (sec.footer, sec.first_page_footer, sec.even_page_footer):
            for el in ft._element.iter(W("r")):
                rpr = el.find(W("rPr"))
                if rpr is None:
                    rpr = el.makeelement(W("rPr"), {})
                    el.insert(0, rpr)
                for x in rpr.findall(W("rFonts")) + rpr.findall(W("sz")) + rpr.findall(W("szCs")):
                    rpr.remove(x)
                rpr.insert(0, rpr.makeelement(W("rFonts"), {W(a): "Times New Roman" for a in ("ascii", "hAnsi", "cs", "eastAsia")}))
                rpr.append(rpr.makeelement(W("sz"), {W("val"): "24"}))
                rpr.append(rpr.makeelement(W("szCs"), {W("val"): "24"}))
    for attempt in range(5):  # файл иногда ненадолго занят (антивирус, индексатор)
        try:
            D.save(OUT)
            break
        except OSError:
            if attempt == 4:
                raise
            time.sleep(1.5)


# ------------------------------------------------------------------ разбор PDF
def norm(s):
    return re.sub(r"\s+", " ", s).strip()


def analyze(splits):
    d = pymupdf.open(PDF)
    pages = [norm(p.get_text()) for p in d]
    toc = {}
    for title, _ in toc_entries():
        t = norm(title)
        toc[title] = next((i + 1 for i in range(2, len(pages)) if t in pages[i]), "")
    fix = None
    for number, b in TABLES:
        rows = b[4]
        dec = dict(splits.get(number, {}))
        move, k = dec.get("move", False), dec.get("k")
        first_len = k or len(rows)
        found = None
        for pi, page in enumerate(d):
            lines = [(l["bbox"][1], norm("".join(s["text"] for s in l["spans"])))
                     for blk in page.get_text("dict")["blocks"] for l in blk.get("lines", [])]
            cap = [y for y, t in lines if t.startswith(f"Таблица {number} –")]
            if not cap:
                continue
            y0 = cap[0]
            nxt = [y for y, t in lines if y > y0 + 1 and (t.startswith("Таблица ") or t.startswith("Продолжение таблицы"))]
            y1 = min(nxt) if nxt else page.rect.height
            ys = sorted(round(it["rect"].y0, 1) for it in page.get_drawings()
                        if it["rect"].height < 1.5 and it["rect"].width > 20 and y0 < it["rect"].y0 < y1)
            merged = [y for i, y in enumerate(ys) if i == 0 or y - ys[i - 1] > 1.5]
            found = (pi, max(0, len(merged) - 1))
            break
        if found is None:
            continue
        on_page = found[1]  # строк (с шапкой) на странице с подписью
        if on_page >= first_len + 1:
            continue
        data_rows = max(0, on_page - 1)
        if number in LONG and data_rows > 0:
            dec["k"] = data_rows                       # длинная: часть строк здесь, остальное — «Продолжение таблицы»
        elif number not in LONG and not dec.get("defer") and can_defer(number):
            dec["defer"] = True                        # сначала текст после таблицы, таблица — следом
        elif not move:
            dec["move"], dec["k"] = True, None         # целиком на следующую страницу
        else:
            dec["k"] = max(1, data_rows)               # не влезает даже на новую страницу
        fix = {number: dec}
        break
    return toc, fix


LO_PYTHON = r"C:\Program Files\LibreOffice\program\python.exe"


def convert():
    """PDF через LibreOffice (UNO): формулам задаются 14 pt и Times New Roman, см. lo_export.py.
    Если LibreOffice завис (бывает после аварийно завершённого экземпляра), процессы и блокировки
    убираются, и экспорт повторяется."""
    lock = OUT.with_name(".~lock." + OUT.name + "#")
    for attempt in range(3):
        lock.unlink(missing_ok=True)
        try:
            # без перехвата вывода: иначе при зависании ожидание каналов не прерывается по таймауту
            before = PDF.stat().st_mtime if PDF.exists() else 0
            r = subprocess.run([LO_PYTHON, str(pathlib.Path(__file__).with_name("lo_export.py")), str(OUT), str(PDF)],
                               stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=60)
            if r.returncode == 0 and PDF.exists() and PDF.stat().st_mtime > before:
                return
            print("экспорт не удался, повтор", flush=True)
        except subprocess.TimeoutExpired:
            print("экспорт завис, перезапуск LibreOffice", flush=True)
        subprocess.run(["powershell", "-NoProfile", "-Command",
                        "Get-CimInstance Win32_Process | Where-Object { $_.CommandLine -like '*lo_profile_qubik*' -or "
                        "$_.CommandLine -like '*lo_export.py*' } | ForEach-Object { Stop-Process -Id $_.ProcessId -Force }; "
                        "Remove-Item \"$env:TEMP\lo_profile_qubik\.lock\" -Force -ErrorAction SilentlyContinue"],
                       capture_output=True, timeout=60)
    raise RuntimeError("не удалось получить PDF")


if __name__ == "__main__":
    toc, splits = {}, {}
    for it in range(1, 16):
        build(toc, splits)
        convert()
        new_toc, fix = analyze(splits)
        print(f"проход {it}: страниц {pymupdf.open(PDF).page_count}, разбиение таблиц {splits}, исправление {fix}", flush=True)
        if fix:
            splits.update(fix)
            toc = new_toc
            continue
        if new_toc == toc:
            break
        toc = new_toc
    print("содержание:", json.dumps(toc, ensure_ascii=False))
