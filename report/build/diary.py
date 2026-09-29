import sys, pathlib
import docx
sys.stdout.reconfigure(encoding="utf-8")
root = pathlib.Path(sys.argv[1])
src = root / "src" / "Дневник_студента_09_535_2_сем_Филадельов .docx"
out_dir = root / "report"; out_dir.mkdir(exist_ok=True)
out = out_dir / "Дневник_студента_09_535_2_сем_Филадельфов.docx"

# номер строки таблицы -> новое содержание (даты и аудиторные занятия не трогаем)
NEW = {
    12: "Реализация модуля загрузки и предварительной обработки табличных наборов данных Iris, Wine, Breast Cancer, Moons и Circles: стратифицированное разбиение, масштабирование признаков, PCA.",
    14: "Реализация базового интерфейса модели для обучения, предсказания и оценки качества (accuracy, precision, recall, F1-score).",
    16: "Реализация классической многослойной нейронной сети Classical MLP с настраиваемым числом и шириной скрытых слоёв.",
    18: "Реализация квантовой нейронной сети Quantum QNN на основе параметризованной квантовой схемы в PennyLane (AngleEmbedding, StronglyEntanglingLayers).",
    20: "Реализация гибридной квантово-классической модели Hybrid QNN: классический энкодер, квантовый слой TorchLayer и выходной слой.",
    22: "Настройка алгоритма обучения: оптимизатор Adam, функция потерь, ранняя остановка по валидационной выборке; исследование влияния скорости обучения и функции активации.",
    24: "Проведение серий экспериментов на пяти наборах данных по 10 запусков: исследование архитектуры и регуляризации перцептрона, сравнение классических моделей с квантовой и гибридной сетями.",
    26: "Анализ результатов экспериментов: accuracy, precision, recall, F1-score и время обучения, проверка значимости различий критерием Уилкоксона.",
    28: "Построение графиков и таблиц результатов, оформление отчёта по практике и текста статьи.",
}

d = docx.Document(src)
t = d.tables[0]
for i, text in NEW.items():
    cell = t.rows[i].cells[1]
    p = cell.paragraphs[0]
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._element.getparent().remove(r._element)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)

# даты — единообразно: ДД.ММ.2026 (в исходнике часть дат была с годом «26», часть — «2026»)
import re
dates = 0
for row in t.rows[1:]:
    cell = row.cells[0]
    m = re.fullmatch(r"\s*(\d{2})\.(\d{2})\.(\d{2}|\d{4})\s*", cell.text)
    if not m:
        continue
    p = cell.paragraphs[0]
    p.runs[0].text = f"{m.group(1)}.{m.group(2)}.2026"
    for r in p.runs[1:]:
        r._element.getparent().remove(r._element)
    dates += 1
print("dates normalized:", dates)

# опечатка в отчестве
fixed = 0
for p in d.paragraphs:
    for r in p.runs:
        if "Геогиевич" in r.text:
            r.text = r.text.replace("Геогиевич", "Георгиевич"); fixed += 1
# --- титульный лист и шапка: поля — подчёркнутые табуляции вместо «____», подписи по центру полей
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
from fix_assignment import run, set_tabs, clear_runs, base_rpr, text_of, W  # noqa: E402

body = d.element.body
sp = body.find(W("sectPr"))
width = int(sp.find(W("pgSz")).get(W("w"))) - int(sp.find(W("pgMar")).get(W("left"))) - int(sp.find(W("pgMar")).get(W("right")))
pars = list(body.iter(W("p")))


def next_text_par(par):
    nxt = par.getnext()
    while nxt is not None and not text_of(nxt).strip():
        nxt = nxt.getnext()
    return nxt


def caption(par, stops_texts):
    rpr = base_rpr(par, False)
    clear_runs(par)
    set_tabs(par, [("center", c) for c, _ in stops_texts])
    for _, txt in stops_texts:
        run(par, rpr, tab=True)
        run(par, rpr, txt)


# «Филадельфов Тихон Георгиевич | 09-535 | подпись»
stud = next(p for p in pars if "Филадельфов Тихон" in text_of(p))
urpr = base_rpr(stud, True)
clear_runs(stud)
set_tabs(stud, [("left", 5200), ("left", 5700), ("center", 6550), ("left", 7400), ("left", 7900), ("right", width)])
run(stud, urpr, "Филадельфов Тихон Георгиевич", underline=True)
run(stud, urpr, tab=True, underline=True)
run(stud, urpr, tab=True)
run(stud, urpr, tab=True, underline=True)
run(stud, urpr, "09-535", underline=True)
run(stud, urpr, tab=True, underline=True)
run(stud, urpr, tab=True)
run(stud, urpr, tab=True, underline=True)
caption(next_text_par(stud), [(2600, "(ФИО студента)"), (6550, "(Группа)"), (8840, "(Подпись)")])

# «должность | ФИО | подпись» для руководителей
for name in ("Васильев А.В.", "Андрианова А.А."):
    par = next(p for p in pars if name in text_of(p) and "_" in text_of(p))
    position = text_of(par).split(name)[0].replace("_", "").strip()
    urpr = base_rpr(par, True)
    clear_runs(par)
    set_tabs(par, [("left", 4200), ("left", 4700), ("center", 6050), ("left", 7400), ("left", 7900), ("right", width)])
    run(par, urpr, position, underline=True)
    run(par, urpr, tab=True, underline=True)
    run(par, urpr, tab=True)
    run(par, urpr, tab=True, underline=True)
    run(par, urpr, name, underline=True)
    run(par, urpr, tab=True, underline=True)
    run(par, urpr, tab=True)
    run(par, urpr, tab=True, underline=True)
    caption(next_text_par(par), [(8840, "(Подпись)")])

# лишний пробел в дате
for par in pars:
    for tx in par.iter(W("t")):
        if tx.text and "февраля  2026" in tx.text:
            tx.text = tx.text.replace("февраля  2026", "февраля 2026")

# шапка второй страницы: «ФИО обучающегося, группа ___Филадельфов Т.Г.____ группа 09-535»
head = next(p for p in pars if text_of(p).startswith("ФИО обучающегося"))
runs = [r for r in head.iter(W("r"))]
label = next(r for r in runs if "ФИО обучающегося" in text_of(r))
urpr = base_rpr(head, True)
for r in runs:
    if r is not label:
        r.getparent().remove(r)
for x in list(head):
    if x.tag == W("proofErr"):
        head.remove(x)
set_tabs(head, [("right", width)])
run(head, base_rpr(head, False), " ")
run(head, urpr, "Филадельфов Т. Г., 09-535", underline=True)
run(head, urpr, tab=True, underline=True)

d.save(out)
print("saved", out, "rows changed:", len(NEW), "typo fixed:", fixed)
for i, r in enumerate(t.rows):
    print(i, r.cells[0].text.strip(), "|", r.cells[1].text.strip()[:110])
